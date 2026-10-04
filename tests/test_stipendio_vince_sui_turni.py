"""«Lo stipendio del mese vince» — fase C2 del piano assistente consulente (04/10/2026).

Turni giornalieri e riga mensile (stipendio da busta paga) dello STESSO
dipendente nello STESSO mese convivono. Per (dipendente, mese) con la riga
mensile il costo viene SOLO dalla busta (ordinario = lordo − extra − chiamata),
i turni danno le ORE (mai sommate a `ore_dichiarate`), non finiscono fra i
«senza costo orario», e le assenze a carico non si sommano (sono in busta).

I test chiamano gli endpoint veri su un client finto che FILTRA davvero le
righe (eq/gte/lte/in_): un mock che restituisce sempre gli stessi dati non
saprebbe dire se una query ha il filtro `mensile=False` — ed e' proprio quel
filtro che impedisce alle copie di saltare il turno del giorno 1.

Valori scelti distinguibili: ogni componente ha un importo che non si ottiene
sommando o sottraendo gli altri per caso.
"""
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import pytest
from fastapi import HTTPException

import services.routers.margini as margini
import services.routers.workspace as workspace
from services.costo_personale_turni import aggrega_per_dipendente_mese


# ---------------------------------------------------------------------------
# Client Supabase finto che filtra
# ---------------------------------------------------------------------------

_DEFAULT = {"mensile": False, "tipo_giorno": "turno", "attiva": True}


class _Query:
    def __init__(self, db, tabella):
        self.db, self.tabella = db, tabella
        self.filtri = []
        self.da_inserire = None
        self.da_aggiornare = None
        self.singola = False

    def select(self, *_a, **_k):
        return self

    def order(self, *_a, **_k):
        return self

    def limit(self, *_a, **_k):
        return self

    def single(self):
        self.singola = True
        return self

    def eq(self, campo, valore):
        self.filtri.append(lambda r: r.get(campo, _DEFAULT.get(campo)) == valore)
        return self

    def in_(self, campo, valori):
        self.filtri.append(lambda r: r.get(campo) in valori)
        return self

    def gte(self, campo, valore):
        self.filtri.append(lambda r: str(r.get(campo)) >= valore)
        return self

    def lte(self, campo, valore):
        self.filtri.append(lambda r: str(r.get(campo)) <= valore)
        return self

    def insert(self, payload):
        self.da_inserire = payload if isinstance(payload, list) else [payload]
        return self

    def update(self, payload):
        self.da_aggiornare = payload
        return self

    def execute(self):
        righe = self.db.setdefault(self.tabella, [])
        if self.da_inserire is not None:
            self.db.setdefault(("inseriti", self.tabella), []).extend(self.da_inserire)
            righe.extend(self.da_inserire)
            return SimpleNamespace(data=list(self.da_inserire))
        trovate = [r for r in righe if all(f(r) for f in self.filtri)]
        if self.da_aggiornare is not None:
            for r in trovate:
                r.update(self.da_aggiornare)
        if self.singola:
            return SimpleNamespace(data=trovate[0] if trovate else None)
        return SimpleNamespace(data=[dict(r) for r in trovate])


class _Client:
    def __init__(self, db):
        self.db = db

    def table(self, nome):
        return _Query(self.db, nome)


def _patch(modulo, db):
    return patch.multiple(
        modulo,
        _resolve_user_from_token=MagicMock(return_value={"id": "user-1"}),
        _get_supabase_client=MagicMock(return_value=_Client(db)),
        _resolve_ristorante_id=MagicMock(return_value="rist-1"),
    )


def _turno(dip, giorno, inizio="09:00", fine="17:00", **kw):
    r = {
        "id": f"t-{dip}-{giorno}", "ristorante_id": "rist-1", "dipendente_id": dip,
        "data_turno": giorno, "mensile": False, "tipo_giorno": "turno",
        "ora_inizio": inizio, "ora_fine": fine,
        "ore_extra": None, "costo_orario": None, "costo_orario_extra": None,
    }
    r.update(kw)
    return r


def _stipendio(dip, mese, lordo, importo_extra=None, importo_chiamata=None,
               ore_dichiarate=0, ore_extra=None):
    return {
        "id": f"m-{dip}-{mese}", "ristorante_id": "rist-1", "dipendente_id": dip,
        "data_turno": f"{mese}-01", "mensile": True, "tipo_giorno": "turno",
        "ora_inizio": "00:00", "ora_fine": "00:00",
        "ore_dichiarate": ore_dichiarate, "ore_extra": ore_extra,
        "lordo_mensile": lordo, "importo_extra": importo_extra,
        "importo_chiamata": importo_chiamata,
        "costo_orario": None, "costo_orario_extra": None,
    }


def _assenza(dip, giorno, tipo, importo):
    return {
        "id": f"a-{dip}-{giorno}", "ristorante_id": "rist-1", "dipendente_id": dip,
        "data_turno": giorno, "mensile": False, "tipo_giorno": tipo,
        "importo_a_carico": importo,
    }


def _costo_mese(turni, anno=2026, mese=6):
    db = {"turni_personale": list(turni)}
    with _patch(margini, db):
        return margini.get_costo_personale_da_turni(anno=anno, mese=mese, authorization="Bearer x")


_DIPENDENTI = [
    {"id": "dip-ada", "ristorante_id": "rist-1", "nome": "Ada", "attivo": True, "costo_orario_default": None},
    {"id": "dip-bruno", "ristorante_id": "rist-1", "nome": "Bruno", "attivo": True, "costo_orario_default": None},
]


def _lista(turni, da="2026-06-01", a="2026-06-30", mensile=None):
    db = {"turni_personale": list(turni), "dipendenti": [dict(d) for d in _DIPENDENTI]}
    with _patch(workspace, db):
        return workspace.ws_personale_list(da=da, a=a, mensile=mensile, authorization="Bearer x")


# Ada: due turni con tariffa (8h + 8h di cui 1 extra) E lo stipendio di giugno.
# Se si pagasse la tariffa: 15×11 + 1×17 = 182. Lo stipendio: 2150 di cui 140
# extra e 310 chiamata -> ordinario 1700.
_ADA_TURNI = [
    _turno("dip-ada", "2026-06-03", costo_orario=11.0, costo_orario_extra=17.0),
    _turno("dip-ada", "2026-06-04", costo_orario=11.0, costo_orario_extra=17.0, ore_extra=1),
]
_ADA_STIPENDIO = _stipendio("dip-ada", "2026-06", 2150.0, importo_extra=140.0,
                            importo_chiamata=310.0, ore_dichiarate=150, ore_extra=9)


# ---------------------------------------------------------------------------
# Margini — GET /api/margini/costo-personale-turni
# ---------------------------------------------------------------------------

class TestMarginiStipendioVince:

    def test_turni_con_tariffa_e_stipendio_costa_solo_lo_stipendio(self):
        res = _costo_mese(_ADA_TURNI + [_ADA_STIPENDIO])
        assert res["costo_dipendenti"] == 1700.0
        assert res["costo_personale_extra"] == 140.0
        assert res["costo_personale_chiamata"] == 310.0
        # Ore dai turni (16, di cui 1 extra), NON 150/9 della busta, e mai 166.
        assert res["ore_totali"] == 16.0
        assert res["ore_extra"] == 1.0
        assert res["n_senza_costo"] == 0
        assert res["n_con_stipendio"] == 1
        # Due turni lavorati; la riga mensile non conta: le ore le danno i turni.
        assert res["n_turni"] == 2

    def test_turni_senza_tariffa_e_stipendio_non_sono_senza_costo(self):
        turni = [
            _turno("dip-ada", "2026-06-03"),
            _turno("dip-ada", "2026-06-05"),
            _stipendio("dip-ada", "2026-06", 1900.0),
            # Bruno senza stipendio e senza tariffa: lui si' resta senza costo.
            _turno("dip-bruno", "2026-06-03", "10:00", "16:00"),
        ]
        res = _costo_mese(turni)
        assert res["n_senza_costo"] == 1
        assert res["costo_dipendenti"] == 1900.0
        assert res["ore_totali"] == 22.0
        assert res["n_con_stipendio"] == 1

    def test_stipendio_senza_turni_prende_le_ore_dichiarate(self):
        res = _costo_mese([_ADA_STIPENDIO])
        assert res["ore_totali"] == 150.0
        assert res["ore_extra"] == 9.0
        assert res["n_turni"] == 1
        assert res["costo_dipendenti"] == 1700.0
        assert res["costo_personale_chiamata"] == 310.0

    def test_la_regola_e_per_dipendente_chi_non_ha_stipendio_paga_a_tariffa(self):
        turni = _ADA_TURNI + [_ADA_STIPENDIO, _turno("dip-bruno", "2026-06-03", "10:00", "16:00", costo_orario=13.0)]
        res = _costo_mese(turni)
        assert res["costo_dipendenti"] == 1700.0 + 78.0
        assert res["costo_personale_extra"] == 140.0
        assert res["n_con_stipendio"] == 1
        assert res["ore_totali"] == 22.0

    def test_assenze_a_carico_ignorate_per_chi_ha_lo_stipendio(self):
        turni = [
            _ADA_STIPENDIO,
            _assenza("dip-ada", "2026-06-10", "ferie", 55.0),
            _assenza("dip-bruno", "2026-06-11", "malattia", 37.0),
        ]
        res = _costo_mese(turni)
        assert res["costo_assenze_a_carico"] == 37.0
        assert res["n_giorni_assenza"] == 2

    def test_chiamata_oltre_il_lordo_non_rende_negativo_l_ordinario(self):
        """Righe scritte prima della guardia: il clamp a 0 resta."""
        res = _costo_mese([_stipendio("dip-ada", "2026-06", 300.0, importo_extra=200.0, importo_chiamata=250.0)])
        assert res["costo_dipendenti"] == 0.0
        assert res["costo_personale_extra"] == 200.0
        assert res["costo_personale_chiamata"] == 250.0

    def test_senza_stipendio_tutto_come_prima(self):
        turni = [
            _turno("dip-bruno", "2026-06-03", costo_orario=10.0, costo_orario_extra=15.0, ore_extra=2),
            _assenza("dip-bruno", "2026-06-04", "ferie", 41.0),
        ]
        res = _costo_mese(turni)
        assert res["costo_dipendenti"] == 60.0
        assert res["costo_personale_extra"] == 30.0
        assert res["costo_personale_chiamata"] == 0.0
        assert res["costo_assenze_a_carico"] == 41.0
        assert res["n_con_stipendio"] == 0


# ---------------------------------------------------------------------------
# Funzione pura
# ---------------------------------------------------------------------------

def test_funzione_pura_separa_i_mesi():
    """Stipendio di giugno: i turni di luglio dello stesso dipendente si pagano a tariffa."""
    from services.fastapi_worker import _ore_turno
    turni = [_ADA_STIPENDIO, _turno("dip-ada", "2026-07-01", costo_orario=12.0)]
    per = aggrega_per_dipendente_mese(turni, _ore_turno)
    assert per[("dip-ada", "2026-06")]["costo_ordinario"] == 1700.0
    assert per[("dip-ada", "2026-06")]["ore"] == 150.0
    assert per[("dip-ada", "2026-07")]["costo_ordinario"] == 96.0
    assert per[("dip-ada", "2026-07")]["con_stipendio"] is False


# ---------------------------------------------------------------------------
# Personale — GET /api/workspace/personale
# ---------------------------------------------------------------------------

class TestPersonaleListStipendioVince:

    def test_per_persona_costo_solo_dallo_stipendio_e_ore_dai_turni(self):
        res = _lista(_ADA_TURNI + [_ADA_STIPENDIO, _assenza("dip-ada", "2026-06-12", "ferie", 55.0)])
        assert res["costo_standard_per_persona"]["Ada"] == 1700.0
        assert res["costo_extra_per_persona"]["Ada"] == 140.0
        assert res["costo_chiamata_per_persona"]["Ada"] == 310.0
        assert res["costo_assenze_per_persona"].get("Ada", 0) == 0
        assert res["monte_ore"]["Ada"] == 16.0
        assert res["ore_standard_per_persona"]["Ada"] == 15.0
        assert res["ore_extra_per_persona"]["Ada"] == 1.0
        assert res["costo_per_persona"]["Ada"] == 2150.0
        assert res["costo_chiamata_totale"] == 310.0
        assert res["costo_totale"] == 2150.0

    def test_periodo_su_due_mesi_applica_la_regola_mese_per_mese(self):
        turni = _ADA_TURNI + [_ADA_STIPENDIO, _turno("dip-ada", "2026-07-01", costo_orario=12.0)]
        res = _lista(turni, da="2026-06-01", a="2026-07-31")
        # giugno: busta (2150 totale); luglio: 8h × 12 = 96.
        assert res["costo_standard_per_persona"]["Ada"] == 1796.0
        assert res["costo_per_persona"]["Ada"] == 2246.0
        assert res["monte_ore"]["Ada"] == 24.0

    def test_stipendio_senza_turni_ore_dalla_busta(self):
        res = _lista([_ADA_STIPENDIO])
        assert res["monte_ore"]["Ada"] == 150.0
        assert res["ore_extra_per_persona"]["Ada"] == 9.0
        assert res["costo_per_persona"]["Ada"] == 2150.0


# ---------------------------------------------------------------------------
# Export Excel
# ---------------------------------------------------------------------------

def _riepilogo(xlsx_bytes):
    from openpyxl import load_workbook
    ws = load_workbook(BytesIO(xlsx_bytes))["Riepilogo"]
    intestazioni = [c.value for c in ws[2]]
    righe = {r[0]: r for r in ws.iter_rows(min_row=3, values_only=True) if r[0]}
    return ws, intestazioni, righe


def test_export_service_colonna_chiamata_prima_del_totale():
    from services.personale_export_service import export_excel_personale_mensile
    xlsx = export_excel_personale_mensile(
        turni=[], dipendenti=[{"id": "d1", "nome": "Ada"}, {"id": "d2", "nome": "Bruno"}],
        mese="2026-06", nome_ristorante="Test",
        ore_standard_per_persona={"Ada": 15.0}, ore_extra_per_persona={"Ada": 1.0},
        costo_standard_per_persona={"Ada": 1700.0, "Bruno": 78.0},
        costo_extra_per_persona={"Ada": 140.0},
        costo_assenze_per_persona={"Bruno": 37.0},
        costo_chiamata_per_persona={"Ada": 310.0},
    )
    ws, intestazioni, righe = _riepilogo(xlsx)
    assert intestazioni == [
        "Dipendente", "Ore std", "Ore extra", "Ore totali",
        "Costo lordo (€)", "Costo ore extra (€)", "Costo assenze (€)", "Costo chiamata (€)", "Totale (€)",
    ]
    assert righe["Ada"][7] == 310.0
    assert righe["Ada"][8] == 2150.0
    assert righe["Bruno"][7] == 0.0
    assert righe["Bruno"][8] == 115.0
    assert righe["TOTALE"][7] == 310.0
    assert righe["TOTALE"][8] == 2265.0
    # Titolo unito su tutte e nove le colonne, la nona larga come le altre.
    assert "A1:I1" in [str(r) for r in ws.merged_cells.ranges]
    assert ws.column_dimensions["I"].width == ws.column_dimensions["E"].width


def test_export_endpoint_stipendio_vince_e_chiamata_nel_totale():
    turni = _ADA_TURNI + [
        _ADA_STIPENDIO,
        _assenza("dip-ada", "2026-06-12", "ferie", 55.0),
        _turno("dip-bruno", "2026-06-03", "10:00", "16:00", costo_orario=13.0),
    ]
    db = {
        "turni_personale": turni,
        "dipendenti": [dict(d) for d in _DIPENDENTI],
        "ristoranti": [{"id": "rist-1", "nome_ristorante": "Trattoria"}],
    }
    with _patch(workspace, db):
        res = workspace.ws_personale_export_mensile(mese="2026-06", authorization="Bearer x")
    _ws, intestazioni, righe = _riepilogo(res.body)
    assert intestazioni[7] == "Costo chiamata (€)"
    assert righe["Ada"][1:9] == (15.0, 1.0, 16.0, 1700.0, 140.0, 0.0, 310.0, 2150.0)
    assert righe["Bruno"][8] == 78.0
    assert righe["TOTALE"][8] == 2228.0


# ---------------------------------------------------------------------------
# Validazioni — POST e PATCH mensile
# ---------------------------------------------------------------------------

def _post_mensile(db=None, **campi):
    base = dict(dipendente_id="dip-ada", mese="2026-06", lordo=1000.0)
    base.update(campi)
    db = db if db is not None else {"turni_personale": [], "dipendenti": [dict(d) for d in _DIPENDENTI]}
    with _patch(workspace, db):
        return workspace.ws_personale_crea_mensile(workspace.TurnoMensileBody(**base), authorization="Bearer x"), db


def _patch_mensile(riga, **campi):
    db = {"turni_personale": [dict(riga)]}
    with _patch(workspace, db):
        workspace.ws_personale_aggiorna_mensile(
            riga["id"], workspace.AggiornaTurnoMensileBody(**campi), authorization="Bearer x"
        )
    return db["turni_personale"][0]


class TestValidazioniMensile:

    def test_post_extra_piu_chiamata_oltre_il_lordo_400(self):
        """600 e 500 stanno ciascuno sotto 1000: solo la SOMMA lo supera."""
        with pytest.raises(HTTPException) as e:
            _post_mensile(importo_extra=600.0, importo_chiamata=500.0)
        assert e.value.status_code == 400
        assert "lordo" in e.value.detail

    def test_post_extra_piu_chiamata_pari_al_lordo_passa(self):
        _res, db = _post_mensile(importo_extra=600.0, importo_chiamata=400.0)
        inserita = db[("inseriti", "turni_personale")][0]
        assert inserita["importo_chiamata"] == 400.0
        assert inserita["ore_dichiarate"] == 0.0

    def test_post_chiamata_negativa_400(self):
        with pytest.raises(HTTPException) as e:
            _post_mensile(importo_chiamata=-20.0)
        assert e.value.status_code == 400
        assert "negativo" in e.value.detail

    def test_post_ore_omesse_valgono_zero_e_lordo_basta(self):
        _res, db = _post_mensile()
        assert db[("inseriti", "turni_personale")][0]["ore_dichiarate"] == 0.0

    def test_post_tutto_a_zero_400(self):
        with pytest.raises(HTTPException) as e:
            _post_mensile(lordo=0.0)
        assert e.value.status_code == 400

    def test_post_ore_extra_senza_ore_totali_400(self):
        """Con ore 0 le ore vengono dai turni: extra dichiarate qui non hanno base."""
        with pytest.raises(HTTPException) as e:
            _post_mensile(ore_extra=4)
        assert e.value.status_code == 400

    def test_post_ore_extra_dentro_le_ore_totali_passa(self):
        _post_mensile(ore_totali=160, ore_extra=12)

    def test_post_doppio_mensile_resta_409(self):
        db = {"turni_personale": [_ADA_STIPENDIO], "dipendenti": [dict(d) for d in _DIPENDENTI]}
        with pytest.raises(HTTPException) as e:
            _post_mensile(db=db)
        assert e.value.status_code == 409

    def test_post_con_turni_nel_mese_passa(self):
        db = {"turni_personale": list(_ADA_TURNI), "dipendenti": [dict(d) for d in _DIPENDENTI]}
        _post_mensile(db=db)
        assert len(db[("inseriti", "turni_personale")]) == 1

    def test_patch_chiamata_si_somma_all_extra_letta_a_db_400(self):
        riga = _stipendio("dip-ada", "2026-06", 1000.0, importo_extra=600.0)
        with pytest.raises(HTTPException) as e:
            _patch_mensile(riga, importo_chiamata=500.0)
        assert e.value.status_code == 400
        assert "lordo" in e.value.detail

    def test_patch_lordo_sotto_extra_piu_chiamata_salvate_400(self):
        riga = _stipendio("dip-ada", "2026-06", 1000.0, importo_extra=300.0, importo_chiamata=350.0)
        with pytest.raises(HTTPException) as e:
            _patch_mensile(riga, lordo=600.0)
        assert e.value.status_code == 400

    def test_patch_chiamata_negativa_400(self):
        riga = _stipendio("dip-ada", "2026-06", 1000.0)
        with pytest.raises(HTTPException) as e:
            _patch_mensile(riga, importo_chiamata=-1.0)
        assert e.value.status_code == 400
        assert "negativo" in e.value.detail

    def test_patch_chiamata_valida_salvata_e_azzerabile(self):
        riga = _stipendio("dip-ada", "2026-06", 1000.0, importo_extra=300.0)
        aggiornata = _patch_mensile(riga, importo_chiamata=250.0)
        assert aggiornata["importo_chiamata"] == 250.0
        azzerata = _patch_mensile(aggiornata, importo_chiamata=None)
        assert azzerata["importo_chiamata"] is None

    def test_patch_ore_extra_con_ore_zero_400(self):
        riga = _stipendio("dip-ada", "2026-06", 1000.0, ore_dichiarate=0)
        with pytest.raises(HTTPException) as e:
            _patch_mensile(riga, ore_extra=5)
        assert e.value.status_code == 400


# ---------------------------------------------------------------------------
# Copie e generazione: la riga mensile (giorno 1) non occupa il giorno
# ---------------------------------------------------------------------------

def _inseriti(db):
    return db.get(("inseriti", "turni_personale"), [])


def test_copia_settimana_non_salta_il_giorno_1_per_la_riga_mensile():
    # 2026-06-01 e' lunedi': la settimana sorgente e' 25-31/5.
    db = {"turni_personale": [
        _turno("dip-ada", "2026-05-25", costo_orario=11.0),
        _stipendio("dip-ada", "2026-06", 2150.0),
    ]}
    with _patch(workspace, db):
        res = workspace.ws_personale_copia_settimana(
            workspace.CopiaSettimanaBody(da="2026-06-01", a="2026-06-07"), authorization="Bearer x"
        )
    assert res["n_copiati"] == 1 and res["n_saltati"] == 0
    assert [r["data_turno"] for r in _inseriti(db)] == ["2026-06-01"]


def test_copia_mese_non_salta_il_giorno_1_per_la_riga_mensile():
    # 4/5/2026 e' il primo lunedi' di maggio, 1/6/2026 il primo di giugno.
    db = {"turni_personale": [
        _turno("dip-ada", "2026-05-04"),
        _stipendio("dip-ada", "2026-06", 2150.0),
        _stipendio("dip-ada", "2026-05", 2150.0),
    ]}
    with _patch(workspace, db):
        res = workspace.ws_personale_copia_mese(
            workspace.CopiaMeseBody(mese="2026-06", dipendente_ids=["dip-ada"]), authorization="Bearer x"
        )
    assert res["n_copiati"] == 1 and res["n_saltati"] == 0
    assert [r["data_turno"] for r in _inseriti(db)] == ["2026-06-01"]


def test_copia_mese_salta_ancora_un_turno_vero_sul_giorno_1():
    db = {"turni_personale": [
        _turno("dip-ada", "2026-05-04"),
        _turno("dip-ada", "2026-06-01", "12:00", "15:00"),
    ]}
    with _patch(workspace, db):
        res = workspace.ws_personale_copia_mese(
            workspace.CopiaMeseBody(mese="2026-06", dipendente_ids=["dip-ada"]), authorization="Bearer x"
        )
    assert res["n_copiati"] == 0 and res["n_saltati"] == 1


def test_genera_da_regole_non_salta_il_giorno_1_per_la_riga_mensile():
    db = {
        "regole_turni_ricorrenti": [{
            "id": "reg1", "ristorante_id": "rist-1", "dipendente_id": "dip-ada",
            "giorno_settimana": 0, "tipo_giorno": "turno", "attiva": True,
            "ora_inizio": "09:00", "ora_fine": "14:00",
            "ora_inizio2": None, "ora_fine2": None, "costo_orario": None,
        }],
        "turni_personale": [_stipendio("dip-ada", "2026-06", 2150.0)],
    }
    with _patch(workspace, db):
        res = workspace.ws_regole_turni_genera(
            workspace.GeneraTurniDaRegoleBody(data_da="2026-06-01", data_a="2026-06-07", dipendente_id=None),
            authorization="Bearer x",
        )
    assert res["n_creati"] == 1 and res["n_saltati"] == 0
    assert [r["data_turno"] for r in _inseriti(db)] == ["2026-06-01"]
